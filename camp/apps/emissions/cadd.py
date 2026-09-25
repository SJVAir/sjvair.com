"""
CARB's California Dairy & Livestock Database (CADD), from its XLSX: each
dairy's location, its herd by year, and its anaerobic digesters.

read() checks the layout and reads everything first; apply() writes it all in
one transaction. A layout this module doesn't know (a CADD version that
renamed or dropped a column) fails in read(), before anything is written.
CARB changes the file's URL with each version.
"""
import zipfile
from collections import Counter
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from django.conf import settings
from django.contrib.gis.geos import Point
from django.db import transaction
from django.utils import timezone

from camp.apps.emissions.models import Dairy, DairyHerd, Digester, animal_units
from camp.apps.regions.models import Region

VERSION = '2.0.0'
URL = 'https://ww2.arb.ca.gov/sites/default/files/2025-10/CADD%20v2.0.0.xlsx'

# Sheet names, compared stripped: CARB's herd sheet is 'Facility Herd Size '.
FACILITIES = 'Facility General Information'
HERDS = 'Facility Herd Size'
DIGESTERS = 'Anaerobic Digesters'
COLUMNS = {
    FACILITIES: (
        'CADDID', 'PlaceID', 'FacilityName', 'Latitude', 'Longitude', 'StreetAddress',
        'City', 'County', 'ZipCode', 'RegionalWaterBoard',
    ),
    HERDS: (
        'CADDID', 'Year', 'MilkCows', 'DryCows', 'OldHeifers', 'YoungHeifers', 'OldCalves',
        'YoungCalves', 'BeefCattle', 'MilkCowsHerdSizeRefCode', 'NonMilkingCattleHerdSizeRefCode',
        'LabeledAsDairy',
    ),
    DIGESTERS: ('CADDID', 'OperationalYear', 'ShutdownYear', 'DataSource'),
}
# Herd sheet column -> DairyHerd field.
HERD_COLUMNS = {
    'MilkCows': 'milk_cows',
    'DryCows': 'dry_cows',
    'OldHeifers': 'old_heifers',
    'YoungHeifers': 'young_heifers',
    'OldCalves': 'old_calves',
    'YoungCalves': 'young_calves',
    'BeefCattle': 'beef_cattle',
}
DAIRY_FIELDS = ('place_id', 'name', 'address', 'point', 'county', 'water_board', 'cadd_version', 'modified')


class CADDFormatError(ValueError):
    """The file isn't the CADD layout this importer knows: a sheet or a column is missing."""


def _blank(value):
    # CARB writes "operating" (ShutdownYear) and some unknowns as the text 'NaN'.
    return value is None or (isinstance(value, str) and value.strip().upper() in ('', 'NAN', 'NA'))


def _int(value):
    return None if _blank(value) else int(float(value))


def _text(value):
    if _blank(value):
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _sheet(workbook, name):
    for title in workbook.sheetnames:
        if title.strip() == name:
            return workbook[title]
    found = ', '.join(repr(title) for title in workbook.sheetnames)
    raise CADDFormatError(f'The CADD file has no {name!r} sheet (it has {found}).')


def read(path):
    """{sheet: [{column: value}, ...]} for the three sheets, skipping rows without a CADDID."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError) as err:
        raise CADDFormatError(f'Not a readable XLSX file: {err}')
    try:
        sheets = {}
        for name, columns in COLUMNS.items():
            rows = _sheet(workbook, name).iter_rows(values_only=True)
            header = [_text(cell) for cell in next(rows, ())]
            missing = [column for column in columns if column not in header]
            if missing:
                raise CADDFormatError(
                    f"The CADD {name!r} sheet has no {', '.join(missing)} column; this importer reads CADD {VERSION}."
                )
            at = {column: header.index(column) for column in columns}
            records = []
            for row in rows:
                record = {column: (row[index] if index < len(row) else None) for column, index in at.items()}
                # Blank rows and footnotes under the data.
                if not _blank(record['CADDID']):
                    records.append(record)
            sheets[name] = records
        return sheets
    finally:
        workbook.close()


@dataclass
class Report:
    dairies_created: int = 0
    dairies_updated: int = 0
    herds: int = 0
    digesters: int = 0
    outside: int = 0
    missing_coordinates: int = 0
    unknown_counties: Counter = field(default_factory=Counter)

    def lines(self):
        lines = [
            f'Dairies: {self.dairies_created:,} added, {self.dairies_updated:,} updated. '
            f'Herd-years: {self.herds:,}. Digesters: {self.digesters:,}.',
            f'Outside the covered counties: {self.outside:,}.',
        ]
        if self.missing_coordinates:
            lines.append(f'Skipped, no coordinates: {self.missing_coordinates:,}.')
        if self.unknown_counties:
            names = ', '.join(f'{name} ({count})' for name, count in sorted(self.unknown_counties.items()))
            lines.append(f'Skipped, unknown county: {names}.')
        return lines


def _dairy_values(row, county, version, report):
    lat, lng = row['Latitude'], row['Longitude']
    if _blank(lat) or _blank(lng):
        report.missing_coordinates += 1
        return None
    return {
        'place_id': _int(row['PlaceID']),
        'name': _text(row['FacilityName'])[:128],
        'address': {'street': _text(row['StreetAddress']), 'city': _text(row['City']), 'zipcode': _text(row['ZipCode'])},
        'point': Point(float(lng), float(lat), srid=4326),
        'county': county,
        'water_board': _text(row['RegionalWaterBoard']),
        'cadd_version': version,
    }


def apply(sheets, version=VERSION):
    """
    Upsert the covered counties' dairies on CADDID and replace their herds
    and digesters, in one transaction. Other counties are counted, not kept;
    a blank county, or a covered one with no Region loaded, is reported.
    """
    report = Report()
    covered = {name.lower() for name in settings.SJVAIR_COUNTIES}
    regions = {region.name.removesuffix(' County').lower(): region for region in Region.objects.counties()}
    rows_by_id = {}
    for row in sheets[FACILITIES]:
        name = _text(row['County'])
        if name and name.lower() not in covered:
            report.outside += 1
            continue
        region = regions.get(name.lower()) if name else None
        if region is None:
            report.unknown_counties[name or '(blank)'] += 1
            continue
        values = _dairy_values(row, region, version, report)
        if values is not None:
            rows_by_id[_int(row['CADDID'])] = values

    with transaction.atomic():
        now = timezone.now()
        existing = {dairy.cadd_id: dairy for dairy in Dairy.objects.filter(cadd_id__in=rows_by_id)}
        new, changed = [], []
        for cadd_id, values in rows_by_id.items():
            dairy = existing.get(cadd_id)
            if dairy is None:
                new.append(Dairy(cadd_id=cadd_id, **values))
                continue
            for name, value in values.items():
                setattr(dairy, name, value)
            dairy.modified = now
            changed.append(dairy)
        Dairy.objects.bulk_create(new, batch_size=500)
        Dairy.objects.bulk_update(changed, DAIRY_FIELDS, batch_size=500)
        report.dairies_created, report.dairies_updated = len(new), len(changed)

        ids = dict(Dairy.objects.filter(cadd_id__in=rows_by_id).values_list('cadd_id', 'pk'))
        DairyHerd.objects.filter(dairy_id__in=ids.values()).delete()
        Digester.objects.filter(dairy_id__in=ids.values()).delete()

        herds = {}
        for row in sheets[HERDS]:
            dairy_id, year = ids.get(_int(row['CADDID'])), _int(row['Year'])
            if dairy_id is None or year is None:
                continue
            counts = {name: _int(row[column]) for column, name in HERD_COLUMNS.items()}
            # One row per dairy and year; a repeated row replaces the earlier one.
            herds[(dairy_id, year)] = DairyHerd(
                dairy_id=dairy_id, year=year, **counts,
                milk_cows_ref_code=_text(row['MilkCowsHerdSizeRefCode']),
                non_milking_ref_code=_text(row['NonMilkingCattleHerdSizeRefCode']),
                labeled_as_dairy=_int(row['LabeledAsDairy']) == 1,
                animal_units=animal_units(counts),
            )
        DairyHerd.objects.bulk_create(herds.values(), batch_size=2000)
        report.herds = len(herds)

        digesters = []
        for row in sheets[DIGESTERS]:
            dairy_id = ids.get(_int(row['CADDID']))
            if dairy_id is None:
                continue
            # A handful of CADD's AgSTAR rows carry no operational year; keep
            # the digester with a blank one rather than dropping it.
            digesters.append(Digester(
                dairy_id=dairy_id, operational_year=_int(row['OperationalYear']),
                shutdown_year=_int(row['ShutdownYear']), source=_text(row['DataSource']),
            ))
        Digester.objects.bulk_create(digesters)
        report.digesters = len(digesters)
    return report
