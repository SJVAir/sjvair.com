"""
CARB's Mandatory GHG Reporting (MRR) annual workbook (https://ww2.arb.ca.gov/mrr-data,
one XLSX per year, released each November). Two sheets matter: '<year> GHG
Data' (one row per reporter, with emitter, fuel-supplier and electricity-
importer figures side by side) and '<year> Emissions by GHG' (tons of CO2,
CH4 and N2O). Only the emitter columns are emissions at a place: suppliers
report fuel they sold. Oil & gas production reports per basin, not per
site. There is no county or coordinate, only city and ZIP.
"""
import os
import re
import tempfile
from collections import Counter

import openpyxl
import requests

from django.db import transaction

from camp.apps.emissions.importers import ghg
from camp.apps.emissions.models import GHGReport
from camp.apps.regions.models import Region

URL = 'https://ww2.arb.ca.gov/sites/default/files/classic/cc/reporting/ghg-rep/reported-data/2024-ghg-emissions-2025-11-04.xlsx'
# ww2.arb.ca.gov answers bare clients with 503; a browser User-Agent gets through.
HEADERS = {'User-Agent': 'Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36 SJVAir'}
DATA_SHEET = '{year} GHG Data'
GAS_SHEET = '{year} Emissions by GHG'
# Column key -> the start of its header, whitespace collapsed. Headers wrap
# and carry qualifiers ("…and CH4 and N2O from Biogenic Fuels"); startswith
# on the collapsed text is the contract, and a renamed column fails read().
COLUMNS = {
    'arb_id': 'ARB ID',
    'name': 'Facility Name',
    'year': 'Report Year',
    'co2e': 'Emitter CO2e from Non-Biogenic Sources',
    'co2e_biogenic': 'Emitter CO2 from Biogenic Fuels',
    'city': 'City',
    'zipcode': 'Zip Code',
    'naics': 'North American Industry Classification System (NAICS)',
    'subparts': 'U.S.EPA/ARB Subparts',
    'sector': 'Industry Sector',
}
GAS_COLUMNS = {'arb_id': 'ARB ID', 'co2': 'CO2', 'ch4': 'CH4', 'n2o': 'N2O'}
HEADER_SCAN_ROWS = 20
BASIN_SECTOR = 'Oil and Gas Production'
_BASIN = re.compile(r'\bbasin\b', re.IGNORECASE)
PROGRAM = GHGReport.Program.MRR


class MRRFormatError(ValueError):
    """The workbook isn't the layout this importer knows: a sheet or a column is missing."""


def download(url=URL):
    """Fetch the workbook to a temp file and return its path (the caller unlinks it). The one network call here."""
    response = requests.get(url, headers=HEADERS, timeout=300)
    response.raise_for_status()
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        tmp.write(response.content)
    return tmp.name


def _text(value):
    return ' '.join(str(value).split()) if value is not None else ''


def _float(value):
    if value is None or value == '':
        return 0.0
    return float(value)


def _zip(value):
    if isinstance(value, (int, float)):
        return f'{int(value):05d}'
    return _text(value)


def _header(sheet, columns, label):
    """(header row number, {key: column index}) for a sheet whose header row starts with 'ARB ID'."""
    # Row numbers are counted here: read-only empty cells carry none.
    for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS), start=1):
        texts = {cell.column: _text(cell.value) for cell in row if cell.value is not None}
        if 'ARB ID' not in texts.values():
            continue
        found = {}
        for key, prefix in columns.items():
            for column, text in texts.items():
                if text.startswith(prefix) and column not in found.values():
                    found[key] = column
                    break
        missing = [columns[key] for key in columns if key not in found]
        if missing:
            raise MRRFormatError(f'{label}: missing column(s) {", ".join(repr(m) for m in missing)}')
        return number, found
    raise MRRFormatError(f'{label}: no header row starting with "ARB ID" in the first {HEADER_SCAN_ROWS} rows')


def read(path, year):
    """(rows, gases) for `year`: rows as dicts keyed by COLUMNS, gases as {arb_id: {co2, ch4, n2o}}."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in (DATA_SHEET.format(year=year), GAS_SHEET.format(year=year)):
        if name not in workbook.sheetnames:
            raise MRRFormatError(f'no sheet named {name!r} (sheets: {", ".join(workbook.sheetnames)})')
    data = workbook[DATA_SHEET.format(year=year)]
    header_row, cols = _header(data, COLUMNS, DATA_SHEET.format(year=year))
    rows = []
    for values in data.iter_rows(min_row=header_row + 1, values_only=True):
        cell = lambda key: values[cols[key] - 1] if cols[key] - 1 < len(values) else None
        arb_id = _text(cell('arb_id'))
        if not arb_id or str(_text(cell('year'))) != str(year):
            continue
        rows.append({
            'arb_id': arb_id, 'name': _text(cell('name')), 'year': year,
            'co2e': _float(cell('co2e')), 'co2e_biogenic': _float(cell('co2e_biogenic')),
            'city': _text(cell('city')), 'zipcode': _zip(cell('zipcode')),
            'naics': _text(cell('naics')).split(' ')[0][:8], 'subparts': _text(cell('subparts')), 'sector': _text(cell('sector')),
        })
    gas = workbook[GAS_SHEET.format(year=year)]
    header_row, cols = _header(gas, GAS_COLUMNS, GAS_SHEET.format(year=year))
    gases = {}
    for values in gas.iter_rows(min_row=header_row + 1, values_only=True):
        cell = lambda key: values[cols[key] - 1] if cols[key] - 1 < len(values) else None
        arb_id = _text(cell('arb_id'))
        if arb_id:
            gases[arb_id] = {key: _float(cell(key)) for key in ('co2', 'ch4', 'n2o')}
    workbook.close()
    return rows, gases


def is_basin_wide(row):
    return row.get('sector') == BASIN_SECTOR and bool(_BASIN.search(row.get('name') or ''))


def valley_zips():
    """{zip5: Region} for every ZIP Region we carry: the Valley filter, since MRR has no county."""
    return {region.external_id: region for region in Region.objects.filter(type=Region.Type.ZIPCODE)}


def apply(rows, gases, year):
    """
    Upsert the Valley emitters for `year` in one transaction and delete that
    year's rows no longer in the file. Keeps a row when its ZIP is one of ours
    and its emitter CO2e is positive; matches by crosswalk, then name + ZIP.
    """
    zips = valley_zips()
    county_of_zip = {}
    counts = Counter()
    seen = set()
    with transaction.atomic():
        for row in rows:
            zip5 = row['zipcode'][:5]
            zip_region = zips.get(zip5)
            if zip_region is None:
                counts['outside'] += 1
                continue
            if row['co2e'] <= 0:
                counts['no_emitter'] += 1
                continue
            facility, method = ghg.resolve(PROGRAM, row['arb_id'], name=row['name'], zipcode=zip5)
            basin = is_basin_wide(row)
            if basin:
                # A basin-wide report is many sites, never one facility's
                # emissions, whatever a name match says; a crosswalk pin
                # stays recorded as that.
                facility, method = None, method if method == GHGReport.MatchMethod.CROSSWALK else ''
            if facility is not None and facility.county_id:
                county = facility.county
            else:
                if zip5 not in county_of_zip:
                    county_of_zip[zip5] = Region.objects.get_county_region(zip_region)
                county = county_of_zip[zip5]
            per_gas = gases.get(row['arb_id'], {})
            GHGReport.objects.update_or_create(
                program=PROGRAM, external_id=row['arb_id'], year=year,
                defaults=dict(
                    facility=facility, match_method=method, county=county,
                    name=row['name'][:128], city=row['city'][:64], zipcode=row['zipcode'][:10], naics=row['naics'],
                    sector=row['sector'][:128], subparts=row['subparts'][:128],
                    co2e=row['co2e'], co2e_biogenic=row['co2e_biogenic'],
                    ch4=per_gas.get('ch4'), n2o=per_gas.get('n2o'),
                    point=None, frs_id='', basin_wide=basin,
                ),
            )
            seen.add(row['arb_id'])
            counts['kept'] += 1
            counts[method or 'unmatched'] += 1
            counts['basin'] += int(basin)
        stale = GHGReport.objects.filter(program=PROGRAM, year=year).exclude(external_id__in=seen)
        counts['deleted'] += stale.count()
        stale.delete()
    return counts


def audit(year, limit=60):
    """Lines for curating the crosswalk: the largest emitters, how each matched, and the same-ZIP candidates otherwise."""
    lines = []
    for report in GHGReport.objects.filter(program=PROGRAM, year=year).select_related('facility').order_by('-co2e')[:limit]:
        if report.basin_wide:
            status = 'basin-wide'
        elif report.facility is not None:
            status = f'{report.match_method} -> {report.facility.name} ({report.facility.county_code}, {report.facility.air_district.external_id}, {report.facility.facid})'
        else:
            status = 'unmatched'
        lines.append(f'{report.external_id:>8} {report.co2e:>14,.0f}  {report.name} [{report.zipcode}]: {status}')
        if report.match_method != GHGReport.MatchMethod.CROSSWALK and not report.basin_wide:
            candidates = sorted(((ghg.similarity(report.name, f.name), f) for f in ghg.candidates_in_zip(report.zipcode)), key=lambda p: -p[0])[:3]
            for score, facility in candidates:
                lines.append(f'{"":>24} {score:.2f} {facility.name} ({facility.county_code}, {facility.air_district.external_id}, {facility.facid})')
    return lines
