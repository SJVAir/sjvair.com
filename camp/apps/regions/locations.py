"""
Import adapters for `regions.Location`: schools and child care facilities.

Each source knows where to download its file, how to parse it into plain
dicts, and which `Location.Type` its rows become. `import_source()` does the
rest: geocoding the rows that arrive without coordinates, resolving the
county and school district, upserting by (source, external_id), and removing
the rows that have disappeared from the source.
"""

import codecs
import csv
import hashlib
import io
import tempfile

from contextlib import contextmanager

import requests

from django.contrib.gis.geos import Point
from django.core.cache import cache
from django.db import transaction
from django.utils import timezone

from camp.apps.regions.models import Location
from camp.utils.geocode import clean_address, resolve


# The eight San Joaquin Valley counties, lowercased and without " County".
SJV_COUNTIES = frozenset([
    'fresno', 'kern', 'kings', 'madera',
    'merced', 'san joaquin', 'stanislaus', 'tulare',
])

# Private schools this small are almost always a family homeschooling under
# the affidavit, not a school with a campus. A blank enrollment doesn't clear
# the bar either: the spec's filter is enrollment >= 6.
MIN_PRIVATE_ENROLLMENT = 6

# CDSS facility types that are child care centers, as FAC_TYPE_DESC spells
# them (the column is truncated at 20 characters, hence "SINGLE CHILD CARE CE"
# for a single child care center). Family child care homes are private
# residences and aren't in this dataset at all.
CHILD_CARE_TYPES = (
    'DAY CARE CENTER',
    'INFANT CENTER',
    'SCHOOL-AGE',
    'SINGLE CHILD CARE',
)
CHILD_CARE_PROGRAM = 'CHILD CARE'

GEOCODE_CACHE_TTL = 60 * 60 * 24 * 30  # 30 days
# An address the geocoder couldn't place is cached too, so a re-import doesn't
# pay for the same failed lookup again -- but for a day rather than a month,
# since the miss may be the geocoder's rather than the address's.
GEOCODE_MISS_TTL = 60 * 60 * 24
GEOCODE_MISS = 'miss'

XLSX_MAGIC = b'PK\x03\x04'

# A source that answers a download with a web page is a bot wall, not data.
HTML_PREFIXES = (b'<html', b'<!doctype')


class DownloadError(Exception):
    """A source file couldn't be fetched."""


# -- Geocoding --

def geocode_cached(address):
    """
    A single-line address → (lat, lng) or None, cached for 30 days. A failed
    lookup is cached as well, for a day (see GEOCODE_MISS_TTL).
    """
    cleaned = clean_address(address)
    if not cleaned:
        return None

    digest = hashlib.sha1(cleaned.encode('utf-8', 'replace')).hexdigest()
    key = f'regions:geocode:{digest}'

    cached = cache.get(key)
    if cached == GEOCODE_MISS:
        return None
    if cached is not None:
        return tuple(cached)

    point = resolve(cleaned)
    if point is None:
        cache.set(key, GEOCODE_MISS, GEOCODE_MISS_TTL)
        return None

    result = (point.y, point.x)
    cache.set(key, result, GEOCODE_CACHE_TTL)
    return result


# -- Row helpers --

def _text(file, encoding):
    """
    Decode a source file. The CDE directory is latin-1 and the data.ca.gov
    exports are UTF-8, so each source says which it is; a UTF-8 source that
    turns out not to be falls back to latin-1 rather than failing the import.
    """
    data = file.read()

    if data.startswith(codecs.BOM_UTF8):
        return data.decode('utf-8-sig')

    if encoding != 'latin-1':
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass

    return data.decode('latin-1')


def _rows(file, delimiter=',', encoding='utf-8'):
    """Yield dicts keyed by the lowercased header, from a binary file object."""
    reader = csv.reader(io.StringIO(_text(file, encoding), newline=''), delimiter=delimiter)

    header = next(reader, None)
    if header is None:
        return

    keys = [_key(name) for name in header]
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        yield dict(zip(keys, [cell.strip() for cell in row]))


def _xlsx_rows(file):
    """Yield dicts keyed by the lowercased header, from an XLSX file object."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError(
            'Reading XLSX requires openpyxl. Export the file to CSV and pass'
            ' it with --path instead.'
        )

    workbook = load_workbook(file, read_only=True, data_only=True)
    rows = workbook[workbook.sheetnames[0]].iter_rows(values_only=True)

    keys = None
    for row in rows:
        values = ['' if value is None else str(value).strip() for value in row]
        if not any(values):
            continue
        if keys is None:
            keys = [_key(value) for value in values]
            continue
        yield dict(zip(keys, values))


def _key(name):
    return str(name or '').strip().lstrip('﻿').strip().lower()


def _get(row, *names, default=''):
    for name in names:
        value = row.get(name)
        if value:
            return value
    return default


def _float(value):
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _int(value):
    try:
        return int(float(str(value).strip().replace(',', '')))
    except (TypeError, ValueError):
        return None


def _in_the_valley(county):
    name = str(county or '').strip().lower()
    if name.endswith(' county'):
        name = name[:-len(' county')]
    return name in SJV_COUNTIES


def _address_string(row):
    parts = [row.get('address') or '', row.get('city') or '']
    zipcode = (row.get('zip') or '').strip()
    parts.append(f'CA {zipcode}'.strip())
    return ', '.join(part for part in parts if part.strip())


# -- Parsers --

def parse_cde_public(file):
    """
    CDE public schools and districts directory (tab-delimited, latin-1).
    District rows carry an empty School; schools carry a 14-digit CDS code
    whose first seven digits identify the district.
    """
    for row in _rows(file, delimiter='\t', encoding='latin-1'):
        name = _get(row, 'school')
        if not name:
            continue  # A district row, not a school.
        if _get(row, 'statustype').lower() != 'active':
            continue
        if _get(row, 'virtual').upper() == 'F':
            continue  # Exclusively virtual: no campus to stand next to.
        if _get(row, 'eilcode').upper() == 'A':
            continue  # Adult education.
        if not _in_the_valley(_get(row, 'county')):
            continue

        cds_code = _get(row, 'cdscode')
        yield {
            'external_id': cds_code,
            'cds_code': cds_code,
            'name': name,
            'address': _get(row, 'street'),
            'city': _get(row, 'city'),
            'zip': _get(row, 'zip'),
            'lat': _float(_get(row, 'latitude')),
            'lng': _float(_get(row, 'longitude')),
            'metadata': {
                'county': _get(row, 'county'),
                'district': _get(row, 'district'),
                'grades': _get(row, 'gsoffered'),
                'soc_type': _get(row, 'soctype'),
                'eil_code': _get(row, 'eilcode'),
                'eil_name': _get(row, 'eilname'),
                'charter': _get(row, 'charter').upper() == 'Y',
                'virtual': _get(row, 'virtual'),
                'status': _get(row, 'statustype'),
            },
        }


def parse_cde_private(file):
    """
    CDE private school affidavit, school level. The columns are renamed most
    years, so everything is read by header name with the known aliases, and
    the file carries no coordinates: these rows are geocoded.
    """
    head = file.read(len(XLSX_MAGIC))
    file.seek(0)
    rows = _xlsx_rows(file) if head == XLSX_MAGIC else _rows(file, encoding='utf-8')

    for row in rows:
        name = _get(row, 'school name', 'schoolname', 'name')
        if not name:
            continue
        if not _in_the_valley(_get(row, 'county', 'county name')):
            continue

        enrollment = _int(_get(row, 'total enrollment', 'enrollment', 'total'))
        if enrollment is None or enrollment < MIN_PRIVATE_ENROLLMENT:
            continue

        address = _get(row, 'street', 'address', 'school street', 'mailing street')
        external_id = _get(row, 'cds code', 'cdscode', 'school code', 'schoolcode',
            'affidavit id', 'affidavitid')
        if not external_id:
            # Some years ship without any stable id. Name + address is the
            # best identity available, and it's stable across re-runs.
            external_id = hashlib.sha1(
                f'{name}|{address}'.lower().encode('utf-8', 'replace')
            ).hexdigest()

        grade_low, grade_high = _grade_span(row)
        yield {
            'external_id': external_id,
            'cds_code': None,
            'name': name,
            'address': address,
            'city': _get(row, 'city', 'school city'),
            'zip': _get(row, 'zip', 'zip code', 'school zip'),
            'lat': _float(_get(row, 'latitude')),
            'lng': _float(_get(row, 'longitude')),
            'metadata': {
                'county': _get(row, 'county', 'county name'),
                'enrollment': enrollment,
                'grade_low': grade_low,
                'grade_high': grade_high,
            },
        }


def _grade_span(row):
    low = _get(row, 'low grade', 'lowgrade', 'grade low')
    high = _get(row, 'high grade', 'highgrade', 'grade high')
    if low or high:
        return low, high

    grades = _get(row, 'grades', 'grade span', 'gsoffered')
    if '-' in grades:
        low, _, high = grades.partition('-')
        return low.strip(), high.strip()
    return grades.strip(), grades.strip()


def parse_cdss_ccl(file):
    """
    CDSS Community Care Licensing facilities, as published on data.ca.gov.
    The file is every licensed facility in the state -- child care, adult and
    senior care, residential -- so the child care centers are picked out by
    PROGRAM_TYPE and FAC_TYPE_DESC. STATUS is a numeric code on the licensed
    layer, so it's recorded rather than filtered on.
    """
    # One site is often licensed as two or three facilities -- a day care
    # center, an infant center and a school-age center at the same name and
    # address, each with its own facility number. Readers see one place, so
    # those merge into one row keyed by the lowest facility number, carrying
    # every number and type and the summed capacity.
    sites = {}
    for row in _rows(file, encoding='utf-8'):
        program = _get(row, 'program_type', 'program').upper()
        if program and program != CHILD_CARE_PROGRAM:
            continue

        facility_type = _get(row, 'fac_type_desc', 'facility type', 'facilitytype', 'type').upper()
        if not facility_type.startswith(CHILD_CARE_TYPES):
            continue

        if not _in_the_valley(_get(row, 'county', 'county name')):
            continue

        name = _get(row, 'name', 'facility name', 'facilityname')
        if not name:
            continue

        facility_number = _get(row, 'fac_nbr', 'facility number', 'facilitynumber', 'facility id')
        address = _get(row, 'res_street_addr', 'facility address', 'address')
        zip_code = _get(row, 'res_zip_code', 'facility zip', 'zip')
        capacity = _int(_get(row, 'capacity', 'facility capacity'))
        key = (name.upper(), address.upper(), zip_code)
        site = sites.get(key)
        if site is None:
            site = sites[key] = {
                'external_id': facility_number,
                'cds_code': None,
                'name': name,
                'address': address,
                'city': _get(row, 'res_city', 'facility city', 'city'),
                'zip': zip_code,
                'lat': _float(_get(row, 'fac_latitude', 'latitude', 'facility latitude', 'y')),
                'lng': _float(_get(row, 'fac_longitude', 'longitude', 'facility longitude', 'x')),
                'metadata': {
                    'county': _get(row, 'county', 'county name'),
                    'facility_type': facility_type,
                    'facility_types': [],
                    'facility_numbers': [],
                    'capacity': 0,
                    'status': _get(row, 'status', 'facility status'),
                    'client_served': _get(row, 'client_served'),
                },
            }
        meta = site['metadata']
        meta['facility_types'].append(facility_type)
        meta['facility_numbers'].append(facility_number)
        meta['capacity'] += capacity or 0
        # The lowest facility number keys the merged row, so it is stable
        # across imports; the day care center is the site's primary type.
        if facility_number < site['external_id']:
            site['external_id'] = facility_number
        if facility_type.startswith('DAY CARE CENTER'):
            meta['facility_type'] = facility_type
    for site in sites.values():
        site['metadata']['facility_types'].sort()
        site['metadata']['facility_numbers'].sort()
        if not site['metadata']['capacity']:
            site['metadata']['capacity'] = None
        yield site


# -- Sources --

SOURCES = {
    'cde-public': {
        'label': 'CDE public schools and districts directory',
        'type': Location.Type.PUBLIC_SCHOOL,
        # cde.ca.gov answers server-side downloads with a JS bot wall (the
        # same one that blocks the OEHHA Prop 65 list), so --path only.
        'url': None,
        'page_url': 'https://www.cde.ca.gov/ds/si/ds/pubschls.asp',
        'parse': parse_cde_public,
    },
    'cde-private': {
        'label': 'CDE private school affidavit (school level)',
        'type': Location.Type.PRIVATE_SCHOOL,
        # The affidavit file is published under a new per-year URL every
        # year, so --path is the route for it too.
        'url': None,
        'page_url': 'https://www.cde.ca.gov/ds/si/ps/',
        'parse': parse_cde_private,
    },
    'cdss-ccl': {
        'label': 'CDSS community care licensing facilities',
        'type': Location.Type.CHILD_CARE,
        'url': None,
        'ckan_dataset': 'community-care-licensing-facilities1',
        'page_url': 'https://data.ca.gov/dataset/community-care-licensing-facilities1',
        'parse': parse_cdss_ccl,
    },
}

CKAN_PACKAGE_URL = 'https://data.ca.gov/api/3/action/package_show'


def has_download(source):
    """Can this source fetch its own file, or does it need --path?"""
    config = SOURCES[source]
    return bool(config.get('url') or config.get('ckan_dataset'))


def _source_url(config):
    if config.get('url'):
        return config['url']

    dataset = config.get('ckan_dataset')
    if dataset:
        try:
            response = requests.get(CKAN_PACKAGE_URL, params={'id': dataset}, timeout=60)
            response.raise_for_status()
            resources = response.json().get('result', {}).get('resources', [])
        except (requests.RequestException, ValueError) as exc:
            raise DownloadError(f'Could not look up {config["label"]}: {exc}')

        for resource in resources:
            if (resource.get('format') or '').upper() == 'CSV' and resource.get('url'):
                return resource['url']

    raise DownloadError(
        f'No download URL for {config["label"]}: get the file from'
        f' {config.get("page_url") or "the source site"} and pass it with --path.'
    )


def _reject_html(config, url, content_type, chunk):
    """
    A 200 that hands back a web page is a bot wall, not the file. Say so
    loudly rather than parsing zero rows and calling it a successful import.
    """
    looks_like_html = (
        'text/html' in content_type
        or bytes(chunk or b'').lstrip()[:9].lower().startswith(HTML_PREFIXES)
    )
    if not looks_like_html:
        return

    raise DownloadError(
        f'{config["label"]} returned a web page instead of a file: the site'
        f' blocks server-side downloads ({url}). Download it in a browser from'
        f' {config.get("page_url") or "the source site"} and pass it with --path.'
    )


@contextmanager
def _open_source(config, path):
    if path:
        with open(path, 'rb') as handle:
            yield handle
        return

    with tempfile.TemporaryFile() as handle:
        url = _source_url(config)
        try:
            response = requests.get(url, timeout=300, stream=True)
            response.raise_for_status()
            content_type = (response.headers.get('Content-Type') or '').lower()
            for index, chunk in enumerate(response.iter_content(chunk_size=1024 * 64)):
                if index == 0:
                    _reject_html(config, url, content_type, chunk)
                handle.write(chunk)
        except requests.RequestException as exc:
            raise DownloadError(f'Could not download {config["label"]}: {exc}')
        handle.seek(0)
        yield handle


# -- Import --

def import_source(source, path=None, geocode=True):
    """
    Import one source into `Location`, returning a dict of counts:
    imported, updated, removed, geocoded, skipped.
    """
    try:
        config = SOURCES[source]
    except KeyError:
        raise ValueError(f'Unknown location source: {source}')

    counts = {'imported': 0, 'updated': 0, 'removed': 0, 'geocoded': 0, 'skipped': 0}

    with _open_source(config, path) as handle:
        rows = [row for row in config['parse'](handle) if row.get('external_id')]

    if not rows:
        # An empty parse means a failed download or a changed format, never
        # that every school in the valley closed. Don't remove anything.
        return counts

    existing = {location.external_id: location
        for location in Location.objects.filter(source=source)}

    # Resolve coordinates first, outside any transaction: geocoding is
    # network work and must not hold a database transaction open.
    seen = set()
    resolved = []
    for row in rows:
        # Every row in the file is "seen", even one we can't place, so the
        # removal step below never deletes a school that's still listed.
        seen.add(row['external_id'])

        point = _point(row.get('lat'), row.get('lng'))
        if point is None:
            current = existing.get(row['external_id'])
            if current is not None:
                # Already imported: keep the coordinates we have.
                resolved.append((row, current.point))
                continue
            if not geocode or not (row.get('address') or '').strip():
                counts['skipped'] += 1
                continue
            result = geocode_cached(_address_string(row))
            if result is None:
                counts['skipped'] += 1
                continue
            counts['geocoded'] += 1
            point = _point(*result)

        resolved.append((row, point))

    with transaction.atomic():
        for row, point in resolved:
            _, created = Location.objects.update_or_create(
                source=source,
                external_id=row['external_id'],
                defaults={
                    'type': config['type'],
                    'name': row['name'][:200],
                    'address': (row.get('address') or '')[:200],
                    'city': (row.get('city') or '')[:100],
                    'zip': (row.get('zip') or '')[:10],
                    'point': point,
                    'county': Location.county_for(point),
                    'district': Location.district_for(point, cds_code=row.get('cds_code')),
                    'metadata': row.get('metadata') or {},
                    'imported_at': timezone.now(),
                },
            )

            counts['imported' if created else 'updated'] += 1

        stale = Location.objects.filter(source=source).exclude(external_id__in=seen)
        counts['removed'] = stale.count()
        stale.delete()

    return counts


def _point(lat, lng):
    if lat is None or lng is None:
        return None
    return Point(float(lng), float(lat), srid=4326)
